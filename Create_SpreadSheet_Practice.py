import openpyxl as xl
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

wb = xl.Workbook()
ws = wb.active
treeData = [["Type", "Leaf Color", "Height"], ["Maple", "Red", 549], ["Oak", "Green", 783], ["Pine", "Green", 1204]]

for row in treeData:
    ws.append(row)

cell_range = ws["A1:C1"]
ft = Font(bold=True)
for row in cell_range:
    for cell in row:
        cell.font = ft

chart = BarChart()
chart.type = "col"
chart.title = "Tree Type"
chart.x_axis.title = 'Tree Type'
chart.y_axis.title = 'height(cms)'
values = Reference(ws,min_col=3, min_row=2, max_row=4, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)
chart.add_data(values)
chart.set_categories(categories)
ws.add_chart(chart, "E1")
wb.save('Tree.xlsx')

